"""Builds the sample brief & evaluation workbook (needs openpyxl):  python3 content/downloads/src/build-xlsx.py"""
import json, pathlib
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

root = pathlib.Path(__file__).resolve().parents[3]
cfg = json.loads((root / "site.config.json").read_text())
GREEN, CREAM, STONE, INK, MUTED = "234840", "FAF7F2", "EFE8DD", "1F1A17", "5E554C"
thin = Side(style="thin", color="D6CCBD")
border = Border(bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = Font(name="Georgia", size=16, bold=True, color=GREEN)
    ws["A2"] = sub
    ws["A2"].font = Font(size=10, italic=True, color=MUTED)
    ws.row_dimensions[1].height = 26

def header(ws, row, cols):
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 30

wb = Workbook()

# 1. Brief
ws = wb.active
ws.title = "1. Coffee brief"
title(ws, "Your coffee brief", "Fill in what you know; “not sure yet” is a useful answer. Share it with your supplier before asking for samples.")
header(ws, 4, ["Item", "Your answer", "Example", "Why it matters"])
rows = [
    ("Product", "", "Green coffee for filter roasting / roasted whole bean for cafés", "Decides which suppliers and formats apply"),
    ("Cup profile", "", "Floral and citrus, light body / chocolate, low acidity for espresso", "Narrows the search faster than region alone"),
    ("Region(s)", "", "Yirgacheffe, Sidama, Guji, Limu, Jimma, Harar, or open", "Ethiopian coffees vary widely by area"),
    ("Process", "", "Washed / natural / open to both", "Shapes the cup; availability differs"),
    ("Quality level", "", "Grade 1 or 2 / minimum cup score under a named protocol / blend component", "Filters what is proposed"),
    ("Quantity per order", "", "e.g. 40 × 60 kg bags", "Decides which coffees and shipping options fit"),
    ("Quantity per year", "", "e.g. 2 containers", "Helps plan across the season"),
    ("Destination", "", "Country and port", "Routes, documents and terms depend on it"),
    ("Timing", "", "Shipment window or crop you are planning for", "Availability shifts through the season"),
    ("Packaging", "", "Jute + hermetic liner / retail packs 250 g", "Confirmed before quotation"),
    ("Certifications / documents", "", "Organic, EUDR information, phytosanitary certificate…", "Must be evidenced, not assumed"),
    ("Preferred terms", "", "FOB Djibouti / CIF your port; payment terms", "Saves a round of questions"),
]
for r, row in enumerate(rows, start=5):
    for c, v in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.alignment = wrap
        cell.border = border
        if c == 1:
            cell.font = Font(bold=True)
        if c == 2:
            cell.fill = PatternFill("solid", fgColor="FFF8E6")
        if c >= 3:
            cell.font = Font(color=MUTED, size=9)
    ws.row_dimensions[r].height = 32
for col, w in zip("ABCD", (24, 38, 46, 38)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"

# 2. Sample log
ws = wb.create_sheet("2. Sample log")
title(ws, "Sample log", "One row per sample received. Keep the lot reference exactly as written on the sample and the offer.")
cols = ["Date received", "Sample type", "Lot reference", "Supplier / exporter", "Region", "Process", "Grade", "Crop year", "Weight received (g)", "Roast date (roasted only)", "Packaging condition", "Notes"]
header(ws, 4, cols)
dv = DataValidation(type="list", formula1='"Type,Offer,Pre-shipment,Arrival,Roasted"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("B5:B200")
for r in range(5, 41):
    for c in range(1, len(cols) + 1):
        ws.cell(row=r, column=c).border = border
for i, w in enumerate((13, 13, 18, 22, 14, 12, 9, 10, 12, 14, 18, 30), start=1):
    ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = w
ws.freeze_panes = "A5"

# 3. Evaluation
ws = wb.create_sheet("3. Evaluation")
title(ws, "Sample evaluation", "Roast and cup every sample the same way. Scores are your own 1–10 impressions, not an official protocol score.")
cols = ["Lot reference", "Moisture %", "Water activity", "Defects (per 350 g)", "Aroma", "Flavour", "Acidity", "Body", "Aftertaste", "Balance", "Overall (1–10)", "Cup consistent across cups?", "Decision", "What you liked", "What didn’t fit"]
header(ws, 4, cols)
dv2 = DataValidation(type="list", formula1='"Approve,Reject,Close - ask for similar,Request pre-shipment sample"', allow_blank=True)
dv3 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
ws.add_data_validation(dv2)
ws.add_data_validation(dv3)
dv2.add("M5:M200")
dv3.add("L5:L200")
for r in range(5, 41):
    for c in range(1, len(cols) + 1):
        ws.cell(row=r, column=c).border = border
for i, w in enumerate((18, 10, 10, 11, 8, 8, 8, 8, 10, 8, 10, 13, 22, 28, 28), start=1):
    ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = w
ws.freeze_panes = "B5"

# 4. Feedback to supplier
ws = wb.create_sheet("4. Feedback email")
title(ws, "Feedback a supplier can act on", "Copy, fill in and send after each sample round.")
text = [
    "Subject: Feedback on sample [LOT REFERENCE]",
    "",
    "Decision: [Approve / Reject / Close, but…]",
    "What we liked: [attributes we want more of]",
    "What didn’t fit: [e.g. too fermented for our customers / acidity higher than our espresso blend needs]",
    "Physical notes: [moisture, defects, packaging, with numbers if available]",
    "Next step: [similar lot / pre-shipment sample / quotation for X bags, shipment window Y]",
    "",
    "When approving, confirm in writing with the lot reference and keep part of the approved sample.",
    "",
    f"Prefer us to handle this? {cfg.get('brand_name','')}: {cfg.get('contact_email','')} · WhatsApp {cfg.get('whatsapp_number','')}",
]
for i, t in enumerate(text, start=4):
    c = ws.cell(row=i, column=1, value=t)
    c.alignment = Alignment(wrap_text=True)
    if t.startswith("Subject") or t.startswith("Prefer"):
        c.font = Font(bold=True, color=GREEN)
ws.column_dimensions["A"].width = 110

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = GREEN

out = root / "content/downloads/sample-brief-and-evaluation-template.xlsx"
wb.save(out)
print("written", out)
